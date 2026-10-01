"""Pruebas de la consolidación de decisiones entre hojas del Excel.

El problema que se resolvió: el archivo le dice al dueño que revise en la hoja
POR REVISAR, pero `--aplicar` solo leía TODOS. Lo que marcaba ahí no se aplicaba
nunca, sin ningún aviso: tres horas de trabajo perdido en silencio.

Ahora se leen todas las hojas y se consolidan. Lo que importa aquí es la REGLA
de prioridad, y sobre todo el caso de contradicción: qué pasa si el mismo
producto está marcado en dos hojas con valores distintos.

Cubren:
  1. Una marca en cualquier hoja cuenta.
  2. El NO explícito gana sobre el SI, en cualquier hoja.
  3. El NO también frena la propuesta automática del reporte.
  4. Sin marcas, no se decide (usa la confianza del reporte).
  5. La columna se busca por nombre aunque se inserten columnas.
"""
from __future__ import annotations

import os
import sys

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, 'herramientas'))

import openpyxl  # noqa: E402

import auditar_catalogo_iva as auditoria  # noqa: E402


def _excel(ruta, hojas):
    """Crea un Excel. `hojas` es {nombre: {id: 'SI'|'NO'}}."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, marcas in hojas.items():
        hoja = libro.create_sheet(nombre)
        hoja.append(['Id BD', 'Nombre del producto', 'APLICAR'])
        for identificador in sorted(marcas):
            hoja.append([identificador, f'PRODUCTO {identificador}',
                         marcas[identificador]])
    libro.save(ruta)
    return ruta


@pytest.fixture
def excel(monkeypatch, tmp_path):
    """Apunta SALIDA_XLSX a un temporal y devuelve una función para crearlo."""
    ruta = str(tmp_path / 'auditoria_catalogo.xlsx')
    monkeypatch.setattr(auditoria, 'SALIDA_XLSX', ruta)
    return lambda hojas: _excel(ruta, hojas)


# ═══════════════════════════════════════════
# 1. Una marca en cualquier hoja cuenta
# ═══════════════════════════════════════════

def test_una_marca_en_por_revisar_se_aplica(excel):
    """El caso que motivó el cambio: se trabaja en POR REVISAR y antes nomogía."""
    excel({'POR REVISAR': {10: 'SI', 11: 'SI'}})

    decisiones = auditoria.decisiones_del_excel()

    assert decisiones == {10: 'SI', 11: 'SI'}


def test_una_marca_en_listos_se_aplica(excel):
    excel({'LISTOS PARA APLICAR': {20: 'SI'}})

    assert auditoria.decisiones_del_excel() == {20: 'SI'}


def test_una_marca_en_todos_se_aplica(excel):
    excel({'TODOS': {30: 'SI'}})

    assert auditoria.decisiones_del_excel() == {30: 'SI'}


def test_marcas_en_tres_hojas_se_consolidan(excel):
    """Un producto puede estar en varias hojas; cada marca cuenta una vez."""
    excel({
        'LISTOS PARA APLICAR': {40: 'SI'},
        'POR REVISAR': {41: 'SI', 42: 'SI'},
        'TODOS': {43: 'SI'},
    })

    decisiones = auditoria.decisiones_del_excel()

    assert decisiones == {40: 'SI', 41: 'SI', 42: 'SI', 43: 'SI'}


def test_el_mismo_producto_marcado_si_en_dos_hojas_gana_si(excel):
    excel({'POR REVISAR': {50: 'SI'}, 'TODOS': {50: 'SI'}})

    assert auditoria.decisiones_del_excel() == {50: 'SI'}


# ═══════════════════════════════════════════
# 2. El NO gana sobre el SI
# ═══════════════════════════════════════════

def test_el_no_de_por_revisar_gana_contra_el_si_de_todos(excel):
    """Contradicción entre hojas: gana el NO.

    Es lo prudente. Si alguien se contradice, lo que sale mal es cambiar la
    tarifa de un producto sin querer, y eso no tiene arreglo; dejarlo como
    estaba, sí.
    """
    excel({'POR REVISAR': {60: 'NO'}, 'TODOS': {60: 'SI'}})

    assert auditoria.decisiones_del_excel() == {60: 'NO'}


def test_el_no_gana_aunque_llegue_despues(excel):
    """El orden de lectura de hojas no cambia el resultado."""
    excel({'LISTOS PARA APLICAR': {61: 'SI'}, 'POR REVISAR': {61: 'NO'}})

    assert auditoria.decisiones_del_excel() == {61: 'NO'}


def test_el_no_frena_a_varios_productos_a_la_vez(excel):
    excel({
        'POR REVISAR': {70: 'SI', 71: 'NO', 72: 'SI'},
        'TODOS': {71: 'SI', 73: 'SI'},
    })

    decisiones = auditoria.decisiones_del_excel()

    assert decisiones == {70: 'SI', 71: 'NO', 72: 'SI', 73: 'SI'}


# ═══════════════════════════════════════════
# 3. El NO frena la propuesta automática
# ═══════════════════════════════════════════

def test_el_no_impide_que_el_reporte_aplique_un_producto_confiable(excel):
    """Este es el motivo del NO: frenar también lo que el reporte propondría.

    Un producto de alta confianza se aplicaría solo. Si el dueño le marcó NO en
    cualquier hoja, esa es su decisión y gana: "a este no lo toques".
    """
    excel({'TODOS': {80: 'NO'}})

    filas = [
        {'id': 80, 'codigo_interno': 'PVP-080', 'aplicar': 'SI'},
        {'id': 81, 'codigo_interno': 'PVP-081', 'aplicar': 'SI'},
    ]
    decisiones = auditoria.decisiones_del_excel()

    candidatos = [f for f in filas
                  if decisiones.get(f['id'], f['aplicar']) == 'SI']

    assert [f['id'] for f in candidatos] == [81], 'El 80 quedó explícitamente afuera'


# ═══════════════════════════════════════════
# 4. Sin marcas, manda el reporte
# ═══════════════════════════════════════════

def test_sin_marcas_no_se_decide_ninguna_hoja(excel):
    excel({'TODOS': {90: ''}})

    decisiones = auditoria.decisiones_del_excel()

    assert 90 not in decisiones, 'Sin marcas debe caer a la confianza del reporte'


def test_un_excel_solo_con_hojas_vacias_no_decide(excel):
    excel({'TODOS': {}, 'POR REVISAR': {}})

    assert auditoria.decisiones_del_excel() == {}


def test_sin_archivo_excel_no_decide(monkeypatch, tmp_path):
    monkeypatch.setattr(auditoria, 'SALIDA_XLSX',
                        str(tmp_path / 'no_existe.xlsx'))

    assert auditoria.decisiones_del_excel() == {}


# ═══════════════════════════════════════════
# 5. Columnas movidas y formas de escribir
# ═══════════════════════════════════════════

def test_encuentra_las_columnas_tras_insertar_una(excel, tmp_path):
    """El dueño puede insertar columnas; la posición no es fija."""
    ruta = str(tmp_path / 'auditoria_catalogo.xlsx')
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'TODOS'
    hoja.append(['Id BD', 'Nombre del producto', 'APLICAR'])
    hoja.append([100, 'PRODUCTO 100', 'SI'])
    hoja.insert_cols(2)
    hoja.cell(row=1, column=2).value = 'Columna insertada'
    hoja.cell(row=2, column=2).value = 'x'
    libro.save(ruta)

    assert auditoria.decisiones_del_excel() == {100: 'SI'}


@pytest.mark.parametrize('texto', ['SI', 'Sí', 'si', 'S', 'YES', '1', 'X'])
def test_todas_las_formas_de_si_cuentan(excel, texto):
    excel({'TODOS': {110: texto}})

    assert auditoria.decisiones_del_excel() == {110: 'SI'}


@pytest.mark.parametrize('texto', ['NO', 'no', 'N', '0'])
def test_todo_lo_que_no_sea_si_cuenta_como_no(excel, texto):
    """Cualquier cosa que no sea un SI claro es un 'no lo toques'."""
    excel({'TODOS': {120: texto}})

    assert auditoria.decisiones_del_excel() == {120: 'NO'}


def test_una_hoja_sin_columna_aplicar_no_rompe(excel):
    """Una hoja rara en el archivo no debe romper la lectura de las demás."""
    ruta = str(excel({'TODOS': {130: 'SI'}}))
    libro = openpyxl.load_workbook(ruta)
    # Se AGREGA una hoja sin la columna APLICAR; TODOS se deja como estaba.
    libro.create_sheet('SIN_APLICAR')['A1'] = 'nada que ver aqui'
    libro.save(ruta)

    assert auditoria.decisiones_del_excel() == {130: 'SI'}


# ═══════════════════════════════════════════
# 6. El detalle para el reporte
# ═══════════════════════════════════════════

def test_el_detalle_dice_de_que_hoja_salio_cada_marca(excel):
    """Para que el comando pueda decir de dónde vino cada decisión."""
    excel({'LISTOS PARA APLICAR': {140: 'SI'},
           'POR REVISAR': {141: 'SI'}})

    decisiones, por_hoja = auditoria.decisiones_del_excel(detallado=True)

    assert decisiones == {140: 'SI', 141: 'SI'}
    assert set(por_hoja) == {'LISTOS PARA APLICAR', 'POR REVISAR'}
    assert por_hoja['POR REVISAR'][141] == 'SI'


def test_una_hoja_inexistente_no_aparece_en_el_detalle(excel):
    excel({'TODOS': {150: 'SI'}})

    _, por_hoja = auditoria.decisiones_del_excel(detallado=True)

    assert 'POR REVISAR' not in por_hoja


def test_el_detalle_de_un_excel_inexistente_no_rompe(monkeypatch, tmp_path):
    monkeypatch.setattr(auditoria, 'SALIDA_XLSX',
                        str(tmp_path / 'nada.xlsx'))

    assert auditoria.decisiones_del_excel(detallado=True) == ({}, {})